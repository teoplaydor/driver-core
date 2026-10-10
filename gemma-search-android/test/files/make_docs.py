"""Documents for TextExtractTest made by real programs (synthetic text, no one's data): LibreOffice writes Word,
OpenDocument, RTF, EPUB and PDF from one page of HTML and Excel / Calc from a CSV; python-pptx a PowerPoint deck (Impress
turns it into .odp); fpdf2 a PDF with an embedded TrueType font (another producer's fonts and spacing); pikepdf the
LibreOffice PDF again with compressed object streams (PDF 1.5).
Also writes pdftotext's text of each PDF (poppler, when installed) to compare words with.

usage: python3 make_docs.py <out_dir>
"""
import os
import shutil
import subprocess
import sys

out = sys.argv[1]
os.makedirs(out, exist_ok=True)

HTML = """<html><head><meta charset="utf-8"><title>Договор</title></head><body>
<h1>Договор поставки № 17/2024</h1>
<p>Поставщик обязуется передать покупателю кофемашины модели «Ромашка-3000» в количестве 12 штук.</p>
<p>Срок поставки — до 15 марта 2025 года. Оплата производится в рублях безналичным переводом.</p>
<table border="1"><tr><td>Товар</td><td>Цена</td></tr><tr><td>Фильтр для воды</td><td>1250 руб.</td></tr></table>
<ul><li>гарантия два года</li><li>доставка курьером до склада</li></ul>
<p>The supplier ships coffee machines to the warehouse in Kazan.</p>
</body></html>
"""
CSV = "Товар,Количество,Склад\nКофемашина,12,Казань\nФильтр для воды,40,Самара\n"


def soffice(src, fmt, outdir, infilter=None):
    cmd = ["soffice", "--headless", "--norestore", "--convert-to", fmt, "--outdir", outdir, src]
    if infilter:
        cmd[2:2] = ["--infilter=" + infilter]
    subprocess.run(cmd, check=True, timeout=180, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)


have_lo = shutil.which("soffice") is not None
work = os.path.join(out, "src")
os.makedirs(work, exist_ok=True)
html = os.path.join(work, "contract.html")
open(html, "w", encoding="utf-8").write(HTML)
csv = os.path.join(work, "stock.csv")
open(csv, "w", encoding="utf-8").write(CSV)

if have_lo:
    soffice(html, 'docx:"MS Word 2007 XML"'.replace('"', ''), out, "HTML (StarWriter)")
    docx = os.path.join(out, "contract.docx")
    soffice(docx, "odt", out)
    soffice(docx, "rtf", out)
    soffice(docx, "pdf", out)
    soffice(docx, "epub", out)
    soffice(csv, "xlsx", out, "CSV:44,34,76,1")
    soffice(csv, "ods", out, "CSV:44,34,76,1")

# a deck: python-pptx, and Impress's .odp of it
from pptx import Presentation

deck = Presentation()
for title, body in [("Отчёт о продажах", "Выручка выросла на 12 процентов"), ("План на весну", "Открыть магазин в Самаре")]:
    s = deck.slides.add_slide(deck.slide_layouts[1])
    s.shapes.title.text = title
    s.placeholders[1].text = body
pptx = os.path.join(out, "deck.pptx")
deck.save(pptx)
if have_lo:
    soffice(pptx, "odp", out)

# fpdf2: DejaVu embedded (a composite font with ToUnicode) and Helvetica (a simple WinAnsi font)
from fpdf import FPDF

font = None
for f in ["/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", "/usr/share/fonts/TTF/DejaVuSans.ttf"]:
    if os.path.exists(f):
        font = f
pdf = FPDF()
pdf.add_page()
if font:
    pdf.add_font("dejavu", "", font)
    pdf.set_font("dejavu", size=14)
    pdf.multi_cell(0, 8, "Квитанция об оплате электроэнергии за октябрь. Сумма к оплате 2 345 рублей. "
                         "Лицевой счёт указан на обороте.", new_x="LMARGIN", new_y="NEXT")
pdf.set_font("helvetica", size=12)
pdf.multi_cell(0, 8, "Electricity bill for October, amount due 2345 roubles.", new_x="LMARGIN", new_y="NEXT")
pdf.output(os.path.join(out, "receipt.pdf"))

# the LibreOffice PDF with object streams
import pikepdf

if have_lo:
    with pikepdf.open(os.path.join(out, "contract.pdf")) as p:
        p.save(os.path.join(out, "objstm.pdf"), object_stream_mode=pikepdf.ObjectStreamMode.generate,
               compress_streams=True)

if shutil.which("pdftotext"):
    for name in os.listdir(out):
        if name.endswith(".pdf"):
            subprocess.run(["pdftotext", "-enc", "UTF-8", os.path.join(out, name), os.path.join(out, name + ".txt")],
                           check=False)
shutil.rmtree(work)
print("ok", sorted(os.listdir(out)))
